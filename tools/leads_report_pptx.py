"""Build standalone PowerPoint decks (ES and FR) from the lead report.

Inputs : data/leads/report_data.json (aggregates) and a texts JSON made by
         `node tools/leads_report_text.js data/leads/informe-leads.html > texts.json`
         (the report's own ES/FR copy with computed numbers, so deck and artifact cannot disagree).
Output : data/leads/informe-leads-es.pptx and informe-leads-fr.pptx (native charts and tables, editable).

Usage  : python -X utf8 tools/leads_report_pptx.py <texts.json>
"""
import json
import math
import sys
from pathlib import Path

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parent.parent
D = json.loads((ROOT / "data" / "leads" / "report_data.json").read_text(encoding="utf-8"))
TX = json.loads(Path(sys.argv[1]).read_text(encoding="utf-8"))
K = TX["K"]

INK, MUTED, ACC, LINE, BG, WHITE = "14202E", "556575", "0B5C7A", "D7DEE5", "F4F6F8", "FFFFFF"
ACC_SOFT, WARN, GOOD = "E1EFF4", "A24A07", "1B7048"
LC = {"es": "C9531A", "fr": "1F63B5", "en": "0F8A67", "ru": "9A4A8C"}
FH, FB = "Georgia", "Calibri"
LANGS = ["es", "en", "fr", "ru"]
SW, SH = 13.333, 7.5


def rgb(h):
    return RGBColor.from_string(h)


def est_height(paras, size, w_in):
    cpl = max(10, int(w_in * 72 / (size * 0.50)))
    lines = sum(max(1, math.ceil(len(p) / cpl)) for p in paras)
    return lines * size * 1.32 / 72 + 0.05 * len(paras)


def add_text(slide, x, y, w, h, paras, size=14, bold=False, color=INK, font=FB, align=None, anchor=None, fit=True, space=4, min_size=9):
    """paras: list of str or (str, dict). dict keys: bold, color, size, bullet."""
    items = [(p, {}) if isinstance(p, str) else p for p in paras]
    if fit:
        while size > min_size and est_height([i[0] + ("  " if i[1].get("bullet") else "") for i in items], size, w - 0.2) > h:
            size -= 0.5
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.05)
    tf.margin_top = tf.margin_bottom = Inches(0.03)
    if anchor:
        tf.vertical_anchor = anchor
    for i, (t, o) in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(space)
        if align:
            p.alignment = align
        r = p.add_run()
        r.text = ("• " if o.get("bullet") else "") + t
        r.font.size = Pt(o.get("size", size))
        r.font.bold = o.get("bold", bold)
        r.font.name = o.get("font", font)
        r.font.color.rgb = rgb(o.get("color", color))
    return tb


def card(slide, x, y, w, h, fill=WHITE, line=LINE):
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    s.adjustments[0] = 0.04
    s.fill.solid()
    s.fill.fore_color.rgb = rgb(fill)
    s.line.color.rgb = rgb(line)
    s.line.width = Pt(0.75)
    s.shadow.inherit = False
    return s


def new_slide(prs, title, kicker=None, lang="es", n=None):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = rgb(BG)
    if kicker:
        add_text(s, 0.55, 0.28, 9, 0.3, [kicker.upper()], size=10, bold=True, color=MUTED, fit=False)
    add_text(s, 0.55, 0.55, 12.2, 0.9, [title], size=26, bold=True, font=FH, color=INK, min_size=18)
    bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.6), Inches(1.42), Inches(0.9), Inches(0.05))
    bar.fill.solid()
    bar.fill.fore_color.rgb = rgb(ACC)
    bar.line.fill.background()
    foot = "Delaguía y Luzón · " + ("Mapa de solicitudes" if lang == "es" else "Cartographie des demandes")
    add_text(s, 0.55, 7.08, 8, 0.3, [foot], size=9, color=MUTED, fit=False)
    if n:
        add_text(s, 12.0, 7.08, 0.8, 0.3, [str(n)], size=9, color=MUTED, align=PP_ALIGN.RIGHT, fit=False)
    return s


def style_chart(chart, legend=True, size=10):
    chart.font.size = Pt(size)
    chart.font.name = FB
    chart.font.color.rgb = rgb(INK)
    chart.has_legend = legend
    if legend:
        chart.legend.position = XL_LEGEND_POSITION.BOTTOM
        chart.legend.include_in_layout = False
        chart.legend.font.size = Pt(size)


def bar_chart(slide, x, y, w, h, cats, values, color, title=None, fmt='0', horizontal=True, size=10):
    cd = CategoryChartData()
    cd.categories = cats
    cd.add_series("n", values)
    gf = slide.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED if horizontal else XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(x), Inches(y), Inches(w), Inches(h), cd)
    ch = gf.chart
    style_chart(ch, legend=False, size=size)
    ch.has_title = bool(title)
    if title:
        ch.chart_title.text_frame.text = title
        ch.chart_title.text_frame.paragraphs[0].runs[0].font.size = Pt(size + 2)
        ch.chart_title.text_frame.paragraphs[0].runs[0].font.bold = True
        ch.chart_title.text_frame.paragraphs[0].runs[0].font.color.rgb = rgb(color)
    pl = ch.plots[0]
    pl.gap_width = 55
    pl.has_data_labels = True
    pl.data_labels.number_format = fmt
    pl.data_labels.number_format_is_linked = False
    pl.data_labels.font.size = Pt(size)
    pl.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END
    ser = pl.series[0]
    ser.format.fill.solid()
    ser.format.fill.fore_color.rgb = rgb(color)
    if horizontal:
        ch.category_axis.reverse_order = True
    ch.category_axis.format.line.color.rgb = rgb(LINE)
    ch.category_axis.tick_labels.font.size = Pt(size)
    ch.value_axis.visible = False
    ch.value_axis.has_major_gridlines = False
    ch.value_axis.maximum_scale = max(values) * 1.25 if values else 1
    return ch


def table(slide, x, y, w, col_w, rows, header=True, size=10, row_h=0.3, hdr_fill=ACC, align_right=()):
    nr, nc = len(rows), len(rows[0])
    gt = slide.shapes.add_table(nr, nc, Inches(x), Inches(y), Inches(w), Inches(row_h * nr))
    tbl = gt.table
    tot = sum(col_w)
    for j, cw in enumerate(col_w):
        tbl.columns[j].width = Emu(int(Inches(w) * cw / tot))
    for i, r in enumerate(rows):
        tbl.rows[i].height = Inches(row_h)
        for j, v in enumerate(r):
            cell = tbl.cell(i, j)
            style = {}
            if isinstance(v, tuple):
                v, style = v
            cell.text = ""
            tf = cell.text_frame
            tf.word_wrap = True
            cell.margin_left = cell.margin_right = Inches(0.06)
            cell.margin_top = cell.margin_bottom = Inches(0.025)
            p = tf.paragraphs[0]
            run = p.add_run()
            run.text = str(v)
            run.font.size = Pt(style.get("size", size))
            run.font.name = FB
            is_h = header and i == 0
            run.font.bold = is_h or style.get("bold", False)
            run.font.color.rgb = rgb(WHITE if is_h else style.get("color", INK))
            if j in align_right:
                p.alignment = PP_ALIGN.RIGHT
            cell.fill.solid()
            cell.fill.fore_color.rgb = rgb(hdr_fill if is_h else style.get("fill", WHITE if i % 2 else "FAFBFC"))
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    return tbl


def fnum(n, lang):
    s = f"{n:,}"
    return s.replace(",", " ") if lang == "fr" else s.replace(",", ".")


def pcts(a, b, lang):
    return f"{round(100 * a / b) if b else 0}" + (" %" if lang == "fr" else " %")


def build(lang):
    T = TX[lang]
    M = T["misc"]
    tl = lambda c: D["tags"][c][lang] if c in D["tags"] else c
    ll = lambda c: D["leaves"][c][lang]
    al = lambda c: D["areas"][c][lang]
    ctry = lambda k: D["country_names"][k][lang] if k in D["country_names"] else ({"US/CA": "EE. UU. / Canadá (+1)" if lang == "es" else "États-Unis / Canada (+1)", "RU/KZ": "Rusia / Kazajistán (+7)" if lang == "es" else "Russie / Kazakhstan (+7)"}.get(k, k))
    fix = lambda s: s.replace("{MULTI}", str(K["multiPct"])).replace("{NAT}", str(K["natPct"]))
    FL = {"es": {l: f"Formulario {l.upper()}" for l in LANGS}, "fr": {l: f"Formulaire {l.upper()}" for l in LANGS}}[lang]
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(SW), Inches(SH)
    n = [0]

    def ns(title, kicker=None):
        n[0] += 1
        return new_slide(prs, title, kicker, lang, n[0] + 1)

    # ---- 1 title
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = rgb(ACC)
    add_text(s, 0.8, 0.9, 11, 0.4, [T["eyebrow"].upper()], size=12, bold=True, color="CFE6EF", fit=False)
    add_text(s, 0.8, 1.5, 11.6, 1.7, [T["title"]], size=40, bold=True, font=FH, color=WHITE, min_size=28)
    add_text(s, 0.8, 3.3, 11.2, 1.5, [T["lead"]], size=16, color="EAF4F8")
    kp = [fnum(D["meta"]["total"], lang), fnum(D["meta"]["coverage"]["text_analysed"], lang), pcts(K["multi"], sum(D["meta"]["text_by_lang"][l] for l in ("es", "en", "fr")), lang),
          f"{round(100 * sum(D['blog_candidate'][l]['n'] for l in ('es', 'fr')) / sum(D['blog_candidate'][l]['of'] for l in ('es', 'fr')))}" + (" %" if lang == "fr" else " %"), f"{K['blog']}" + (" %" if lang == "fr" else " %")]
    for i, (v, lab) in enumerate(zip(kp, T["kpi"])):
        x = 0.8 + i * 2.4
        card(s, x, 5.3, 2.25, 1.4, fill="0E4A63", line="2C7A99")
        add_text(s, x + 0.1, 5.38, 2.05, 0.6, [v], size=26, bold=True, font=FH, color=WHITE, fit=False)
        add_text(s, x + 0.1, 6.0, 2.05, 0.65, [lab], size=11, color="CFE6EF")

    # ---- 2-3 essentials
    f = T["find"]
    for part, items in enumerate((f[:4], f[4:])):
        s = ns(T["h"]["esencial"] + (f" ({part + 1}/2)" if True else ""), T["eyebrow"])
        h_each = 1.2 if part == 0 else 1.55
        for i, (head, body) in enumerate(items):
            y = 1.7 + i * (h_each + 0.1)
            card(s, 0.6, y, 12.1, h_each)
            add_text(s, 0.8, y + 0.08, 11.7, 0.4, [head], size=16, bold=True, color=ACC, min_size=13)
            add_text(s, 0.8, y + 0.48, 11.7, h_each - 0.52, [body], size=12.5, color=INK)

    # ---- volume
    s = ns(T["h"]["volumen"], T["eyebrow"])
    qs = sorted({q for l in LANGS for q in D["quarters"][l]})
    cd = CategoryChartData()
    cd.categories = [f"T{q[-1]} {q[2:4]}" for q in qs]
    for l in ("es", "fr", "en", "ru"):
        cd.add_series(FL[l], [D["quarters"][l].get(q, 0) for q in qs])
    gf = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_STACKED, Inches(0.6), Inches(1.7), Inches(8.2), Inches(5.2), cd)
    ch = gf.chart
    style_chart(ch, True, 10)
    ch.plots[0].gap_width = 45
    for ser, l in zip(ch.plots[0].series, ("es", "fr", "en", "ru")):
        ser.format.fill.solid()
        ser.format.fill.fore_color.rgb = rgb(LC[l])
    ch.value_axis.has_major_gridlines = True
    ch.value_axis.major_gridlines.format.line.color.rgb = rgb(LINE)
    ch.value_axis.format.line.fill.background()
    card(s, 9.0, 1.7, 3.75, 5.2)
    add_text(s, 9.15, 1.85, 3.45, 4.9, [T["i_volumen"]], size=13)

    # ---- type of lead
    s = ns(T["h"]["tipo"], T["eyebrow"])
    order = ["urgente", "presupuesto", "cita", "informacion", "caso_detallado", "otro"]
    cols = {"urgente": WARN, "presupuesto": LC["fr"], "cita": LC["en"], "informacion": "8AA5B8", "caso_detallado": LC["ru"], "otro": "CBD9E3"}
    cd = CategoryChartData()
    langs3 = ["es", "fr", "en"]
    cd.categories = [FL[l] for l in langs3]
    for k in order:
        vals = []
        for l in langs3:
            nat = D["nature"][l]
            base = sum(nat.get(o, 0) for o in order) or 1
            vals.append(round(100 * nat.get(k, 0) / base, 1))
        cd.add_series(T["natK"][k], vals)
    gf = s.shapes.add_chart(XL_CHART_TYPE.BAR_STACKED_100, Inches(0.6), Inches(1.7), Inches(7.6), Inches(3.4), cd)
    ch = gf.chart
    style_chart(ch, True, 10)
    ch.plots[0].gap_width = 40
    ch.category_axis.reverse_order = True
    ch.value_axis.visible = False
    ch.value_axis.has_major_gridlines = False
    for ser, k in zip(ch.plots[0].series, order):
        ser.format.fill.solid()
        ser.format.fill.fore_color.rgb = rgb(cols[k])
    ch.has_title = True
    ch.chart_title.text_frame.text = M["nature"]
    ch.chart_title.text_frame.paragraphs[0].runs[0].font.size = Pt(12)
    ch.chart_title.text_frame.paragraphs[0].runs[0].font.bold = True
    rows = [["", M["client"], M["fP"], M["fA"], M["fU"], M["fI"]]]
    for l in langs3:
        c = D["client"][l]
        tn = D["meta"]["text_by_lang"][l]
        fl = D["flags"][l]
        rows.append([FL[l], f"{round(100 * c['p'] / (c['p'] + c['c']))} · {round(100 * c['c'] / (c['p'] + c['c']))} %"] + [f"{round(100 * fl[k] / tn)} %" for k in ("p", "a", "u", "i")])
    c = D["client"]["ru"]
    rows.append([FL["ru"], f"{round(100 * c['p'] / (c['p'] + c['c']))} · {round(100 * c['c'] / (c['p'] + c['c']))} %", "n/d", "n/d", "n/d", "n/d"])
    add_text(s, 0.6, 5.2, 7.6, 0.3, [M["flagsT"]], size=12, bold=True, color=ACC, fit=False)
    table(s, 0.6, 5.5, 7.6, [1.6, 1.6, 1.2, 1.2, 1.2, 1.5], rows, size=9, row_h=0.27, align_right=(1, 2, 3, 4, 5))
    card(s, 8.5, 1.7, 4.25, 5.2)
    add_text(s, 8.65, 1.85, 3.95, 4.9, [T["i_tipo"], (M["medLen"] + f": ES {D['msg_len_median']['es']}, FR {D['msg_len_median']['fr']} " + M["chars"], {"color": MUTED, "size": 12})], size=13)

    # ---- origin
    s = ns(T["h"]["origen"], T["eyebrow"])
    for i, l in enumerate(("es", "fr", "en")):
        cs = D["country_split"][l]
        items = D["country_by_lang"][l][:7]
        bar_chart(s, 0.6 + i * 4.1, 1.7, 3.95, 3.4, [ctry(k) for k, _ in items], [v for _, v in items], LC[l], title=FL[l], size=9)
        add_text(s, 0.6 + i * 4.1, 5.1, 3.95, 0.6, [f"{pcts(cs['spain'], cs['with_country'], lang)} {M['inSpain']} · {pcts(cs['other'], cs['with_country'], lang)} {M['outSpain']} · {cs['none']} {M['unknown']}"], size=10, color=MUTED)
    add_text(s, 0.6, 5.75, 12.1, 1.3, [T["i_origen"], (M["cNote"] + " " + M["traffic"], {"color": MUTED, "size": 10})], size=11)

    # ---- nationality
    s = ns(M["natT"], T["h"]["origen"])
    for i, l in enumerate(("fr", "es", "en")):
        items = D["nat_by_lang"][l][:8]
        bar_chart(s, 0.6 + i * 4.1, 1.7, 3.95, 3.9, [ctry(k) for k, _ in items], [v for _, v in items], LC[l], title=FL[l], size=9)
        add_text(s, 0.6 + i * 4.1, 5.6, 3.95, 0.4, [f"{D['nat_n'][l]} / {D['meta']['text_by_lang'][l]} " + ("mensajes declaran nacionalidad" if lang == "es" else "messages déclarent une nationalité")], size=10, color=MUTED)
    fam = " · ".join(f"{ctry(k)} {v}" for k, v in D["fam_nat"][:8])
    add_text(s, 0.6, 6.1, 12.1, 0.9, [(M["famNat"] + ": " + fam, {"size": 11})], size=11)

    # ---- areas
    s = ns(T["h"]["areas"], T["eyebrow"])
    AR = ["EXT", "FIS", "INM", "SUC", "MER", "LAB", "OTR"]

    def cell(obj, a):
        tot = sum(obj.values()) or 1
        v = obj.get(a, 0)
        return f"{v} ({round(100 * v / tot)} %)"

    rows = [["", f"ES {M['colDD']}", f"ES {M['colTX']}", f"FR {M['colDD']}", f"FR {M['colTX']}", f"EN {M['colDD']}", f"RU {M['colDD']}"]]
    for a in AR:
        rows.append([M["other"] if a == "OTR" else al(a), cell(D["dropdown_area"]["es"], a), cell(D["text_area"]["es"], a), cell(D["dropdown_area"]["fr"], a), cell(D["text_area"]["fr"], a), cell(D["dropdown_area"]["en"], a), cell(D["dropdown_area"]["ru"], a)])
    table(s, 0.6, 1.75, 12.1, [3.4, 1.4, 1.4, 1.4, 1.4, 1.4, 1.4], rows, size=11, row_h=0.42, align_right=(1, 2, 3, 4, 5, 6))
    add_text(s, 0.6, 5.3, 12.1, 1.6, [T["i_areas"]], size=13)

    # ---- by language
    s = ns(T["h"]["idiomas"] + " (1/2)", T["eyebrow"])
    add_text(s, 0.6, 1.55, 12.1, 0.55, [T["i_idiomas"]], size=11, color=MUTED)
    for i, l in enumerate(("es", "fr")):
        x = 0.6 + i * 6.15
        card(s, x, 2.2, 5.95, 4.75)
        add_text(s, x + 0.15, 2.28, 5.65, 0.45, [T["cards"][l]["h"]], size=14, bold=True, color=LC[l], min_size=11)
        add_text(s, x + 0.15, 2.75, 5.65, 1.55, [T["cards"][l]["p"]], size=11)
        nt = D["meta"]["text_by_lang"][l]
        top = sorted(((c, v.get(l, 0)) for c, v in D["tag_langs"].items() if not c.startswith("x_") and c not in ("ext_general", "fis_general", "inm_general")), key=lambda x: -x[1])[:6]
        bar_chart(s, x + 0.1, 4.3, 3.6, 2.6, [tl(c) for c, _ in top], [round(100 * v / nt, 1) for _, v in top], LC[l], fmt='0"%"', size=8)
        lf = D["lift"][l][:5]
        add_text(s, x + 3.75, 4.35, 2.1, 2.55, [(M["lift"], {"bold": True, "color": ACC, "size": 10})] + [(f"{tl(r[0])} ×{round(r[4], 1) if r[4] < 10 else round(r[4])}", {"size": 9}) for r in lf], size=9, space=3)
    s = ns(T["h"]["idiomas"] + " (2/2)", T["eyebrow"])
    for i, l in enumerate(("en", "ru")):
        x = 0.6 + i * 6.15
        card(s, x, 1.65, 5.95, 2.35)
        add_text(s, x + 0.15, 1.72, 5.65, 0.45, [T["cards"][l]["h"]], size=14, bold=True, color=LC[l], min_size=11)
        add_text(s, x + 0.15, 2.2, 5.65, 1.75, [T["cards"][l]["p"]], size=12)
    per = D["period_share"]
    tops = sorted(per[D["periods"][2]]["tags"], key=lambda k: -max(per[p]["tags"].get(k, 0) for p in D["periods"]))[:8]
    rows = [[M["pTitle"]] + [f"{p.replace('..', ' → ')} (n = {per[p]['n']})" for p in D["periods"]]]
    for k in tops:
        rows.append([tl(k)] + [f"{per[p]['tags'].get(k, 0):.1f} %".replace(".", ",") for p in D["periods"]])
    table(s, 0.6, 4.2, 12.1, [5, 2.4, 2.4, 2.4], rows, size=10, row_h=0.28, align_right=(1, 2, 3))
    add_text(s, 0.6, 6.78, 12, 0.25, [M["pNote"]], size=9, color=MUTED, fit=False)

    # ---- tree
    order = {"SUC": ["SUC-0", "SUC-1", "SUC-2"], "INM": ["INM-1"], "LAB": ["LAB-1", "LAB-2", "LAB-3", "LAB-4"], "MER": ["MER-1", "MER-2", "MER-3", "MER-4"], "FIS": ["FIS-1", "FIS-2", "FIS-3"], "EXT": ["EXT-1", "EXT-2", "EXT-3", "EXT-4"]}
    groups = [("EXT",), ("FIS", "INM"), ("SUC", "MER"), ("LAB",)]
    for gi, g in enumerate(groups):
        s = ns(T["h"]["arbol"] + f" ({gi + 1}/{len(groups)})", T["eyebrow"])
        y = 1.6
        if gi == 0:
            add_text(s, 0.6, y, 12.1, 0.5, [T["i_arbol"]], size=10, color=MUTED)
            y += 0.55
        for a in g:
            add_text(s, 0.6, y, 12.1, 0.85, [(al(a) + " · " + T["areaNotes"][a], {"size": 11})], size=11, color=INK)
            y += 0.95 if len(T["areaNotes"][a]) > 190 else 0.6
            rows = [[M["colLeaf"], M["colN"], M["colPrim"], M["colLang"], M["colSub"]]]
            for k in order[a]:
                st = D["leaf_stats"].get(k, {"n": 0, "primary": 0, "langs": {}, "subtopics": []})
                lf = D["leaves"][k]
                prop = not lf["url"]
                subs = " · ".join(f"{tl(c)} {v}" for c, v in st["subtopics"][:4])
                rows.append([(lf[lang] + ("  [" + M["proposed"] + "]" if prop else ""), {"color": WARN if prop else INK, "bold": prop}), str(st["n"]), str(st["primary"]), f"{st['langs'].get('es', 0)} · {st['langs'].get('en', 0)} · {st['langs'].get('fr', 0)}", (subs or "-", {"size": 10})])
            table(s, 0.6, y, 12.1, [3.5, 0.9, 1.0, 1.1, 5.6], rows, size=11, row_h=0.5, align_right=(1, 2, 3))
            y += 0.5 * len(rows) + 0.2

    # ---- combos
    s = ns(T["h"]["combos"], T["eyebrow"])
    add_text(s, 0.6, 1.55, 12.1, 0.55, [T["i_combos"]], size=11, color=MUTED)
    rows = [[M["colPair"], M["colPairN"], M["pkg"]]]
    for a, b, c in D["combos"][:12]:
        rows.append([f"{tl(a)} + {tl(b)}", str(c), T["pk"].get(f"{a}|{b}", "")])
    table(s, 0.6, 2.2, 12.1, [8, 1.2, 3], rows, size=10, row_h=0.36, align_right=(1,))

    # ---- form: 4 steps
    s = ns(T["h"]["form"], T["eyebrow"])
    add_text(s, 0.6, 1.55, 12.1, 0.7, [T["i_form"]], size=11, color=MUTED)
    for i, (k, head, bullets) in enumerate(T["steps"]):
        x = 0.6 + i * 3.05
        card(s, x, 2.35, 2.9, 4.5)
        add_text(s, x + 0.12, 2.42, 2.6, 0.3, [k], size=11, bold=True, color=ACC, fit=False)
        add_text(s, x + 0.12, 2.75, 2.6, 0.7, [head], size=16, bold=True, font=FH, min_size=12)
        add_text(s, x + 0.12, 3.5, 2.65, 3.25, [(fix(b), {"bullet": True}) for b in bullets], size=11)

    # ---- routes
    keys = sorted(D["intents"]["stats"], key=lambda k: -D["intents"]["stats"][k]["n"])
    for gi in range(0, len(keys), 3):
        s = ns(M["routesT"] if "routesT" in M else T["h"]["form"], T["h"]["form"])
        for j, k in enumerate(keys[gi:gi + 3]):
            st = D["intents"]["stats"][k]
            name, qs = T["ints"][k]
            x = 0.6 + j * 4.1
            card(s, x, 1.65, 3.95, 5.3)
            add_text(s, x + 0.12, 1.7, 3.7, 0.55, [name], size=13, bold=True, color=ACC, min_size=10)
            add_text(s, x + 0.12, 2.25, 3.7, 0.3, [f"{st['n']} {M['colShare'].lower()} · {pcts(st['bc'], st['bc_of'], lang)} {M['info']}"], size=10, color=MUTED, fit=False)
            add_text(s, x + 0.12, 2.6, 3.7, 3.05, [(f"{i + 1}. {q}", {}) for i, q in enumerate(qs)], size=11.5, space=4)
            arts = D["suggest"][k][lang][:2]
            add_text(s, x + 0.12, 5.7, 3.7, 1.2, [("Guías hoy (selección automática, a curar)" if lang == "es" else "Guides aujourd'hui (sélection automatique, à trier)", {"bold": True, "color": ACC, "size": 9})] + [(a["title"], {"size": 9}) for a in arts], size=9, space=2)

    # ---- blog
    s = ns(T["h"]["blog"], T["eyebrow"])
    add_text(s, 0.6, 1.55, 12.1, 0.7, [T["i_blog"]], size=10, color=MUTED)
    cand = []
    for k, lg in D["tag_langs"].items():
        if k.startswith("x_") or k in ("ext_general", "fis_general", "inm_general"):
            continue
        ps = D["posts_by_tag"][k]
        tot = sum(lg.values())
        gaps = [l for l in ("es", "fr") if lg.get(l, 0) >= 8 and ps[l] <= 1]
        if tot >= 15:
            cand.append((bool(gaps), tot, k, lg, ps, gaps))
    cand.sort(key=lambda r: (not r[0], -r[1]))
    rows = [[M["colSubB"], M["colLeads"], M["colPosts"], ""]]
    for _, _, k, lg, ps, gaps in cand[:14]:
        rows.append([tl(k), f"{lg.get('es', 0)} · {lg.get('fr', 0)} · {lg.get('en', 0)}", f"{ps['es']} · {ps['en']} · {ps['fr']}", ((M["gap"] + " " + ", ".join(x.upper() for x in gaps)), {"color": WARN, "bold": True}) if gaps else (M["ok"], {"color": GOOD})])
    table(s, 0.6, 2.25, 12.1, [5.6, 2, 2, 2.5], rows, size=10, row_h=0.285, align_right=(1, 2))
    add_text(s, 0.6, 6.72, 12.1, 0.3, [M["gapN"] + " " + M["enNote"]], size=9, color=MUTED, fit=False)

    # ---- recommendations
    s = ns(T["h"]["recom"], T["eyebrow"])
    for i, (head, body) in enumerate(T["recs"]):
        y = 1.65 + i * 0.76
        card(s, 0.6, y, 12.1, 0.68)
        add_text(s, 0.75, y + 0.02, 3.9, 0.64, [fix(head)], size=12, bold=True, color=ACC, anchor=MSO_ANCHOR.MIDDLE, min_size=9)
        add_text(s, 4.7, y + 0.02, 7.9, 0.64, [fix(body)], size=10.5, anchor=MSO_ANCHOR.MIDDLE, min_size=8)

    # ---- method
    s = ns(T["h"]["metodo"], T["eyebrow"])
    add_text(s, 0.6, 1.65, 12.1, 4.3, [(fix(m), {"bullet": True}) for m in T["metodo"]], size=13, space=8)
    add_text(s, 0.6, 6.0, 12.1, 0.9, [(M["files"] + ": data/leads/leads_es.csv, leads_en.csv, leads_fr.csv, leads_ru.csv, topic_article_map.csv, leads_stats.json", {"size": 11, "color": MUTED}), (M["foot"], {"size": 10, "color": MUTED})], size=11)

    out = ROOT / "data" / "leads" / f"informe-leads-{lang}.pptx"
    prs.save(out)
    print(out.name, len(prs.slides), "slides", out.stat().st_size, "bytes")


if __name__ == "__main__":
    build("es")
    build("fr")
